"""Document ingestion: validate -> extract -> clean -> chunk -> metadata.

The HRM never memorizes uploads. Chunks are indexed by the existing RAG layer.
"""

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any

ALLOWED_SUFFIXES = {".pdf", ".docx", ".txt", ".md", ".pptx"}
MAX_BYTES = 8 * 1024 * 1024


class DocumentError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def validate_file(path: str | Path, *, max_bytes: int = MAX_BYTES) -> Path:
    path = Path(path)
    if not path.exists() or not path.is_file():
        raise DocumentError("FILE_NOT_FOUND", "Upload not found")
    if path.suffix.lower() not in ALLOWED_SUFFIXES:
        raise DocumentError("UNSUPPORTED_FILE", f"Unsupported type {path.suffix}")
    size = path.stat().st_size
    if size <= 0 or size > max_bytes:
        raise DocumentError("FILE_TOO_LARGE", "File empty or exceeds size limit")
    return path


def extract_text(path: str | Path) -> str:
    path = validate_file(path)
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".pptx":
        return _extract_pptx(path)
    raise DocumentError("UNSUPPORTED_FILE", f"Unsupported type {suffix}")


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:
        try:
            from PyPDF2 import PdfReader  # type: ignore
        except ImportError as exc:
            raise DocumentError("NOT_IMPLEMENTED", "PDF extraction requires pypdf (not installed in this environment)") from exc
    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n\n".join(pages)


def _extract_docx(path: Path) -> str:
    try:
        import docx  # type: ignore
    except ImportError as exc:
        raise DocumentError("NOT_IMPLEMENTED", "DOCX extraction requires python-docx") from exc
    document = docx.Document(str(path))
    return "\n".join(p.text for p in document.paragraphs)


def _extract_pptx(path: Path) -> str:
    try:
        from pptx import Presentation  # type: ignore
    except ImportError as exc:
        raise DocumentError("NOT_IMPLEMENTED", "PPTX extraction requires python-pptx") from exc
    pres = Presentation(str(path))
    texts = []
    for slide in pres.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                texts.append(shape.text)
    return "\n".join(texts)


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


_SECTION_RE = re.compile(r"(?im)^(unit\s+\d+|chapter\s+\d+|module\s+\d+|section\s+\d+)[^\n]{0,80}$")


def detect_sections(text: str) -> list[dict[str, str]]:
    lines = text.splitlines()
    sections: list[dict[str, str]] = []
    current = "preamble"
    buf: list[str] = []
    for line in lines:
        if _SECTION_RE.match(line.strip()):
            if buf:
                sections.append({"section": current, "content": "\n".join(buf).strip()})
            current = line.strip()
            buf = []
        else:
            buf.append(line)
    if buf:
        sections.append({"section": current, "content": "\n".join(buf).strip()})
    return [s for s in sections if s["content"]]


def chunk_text(text: str, *, max_chars: int = 700, overlap: int = 80) -> list[str]:
    text = clean_text(text)
    if not text:
        return []
    chunks = []
    i = 0
    while i < len(text):
        chunk = text[i : i + max_chars]
        chunks.append(chunk.strip())
        i += max_chars - overlap
    return [c for c in chunks if c]


def extract_metadata(path: Path, text: str) -> dict[str, Any]:
    digest = hashlib.sha256(text.encode("utf-8", errors="replace")).hexdigest()
    return {
        "filename": path.name,
        "suffix": path.suffix.lower(),
        "bytes": path.stat().st_size,
        "sha256": digest,
        "chars": len(text),
    }


def ingest_document(path: str | Path) -> dict[str, Any]:
    path = validate_file(path)
    raw = extract_text(path)
    text = clean_text(raw)
    if not text:
        raise DocumentError("EMPTY_DOCUMENT", "No extractable text")
    sections = detect_sections(text)
    chunks = []
    if sections:
        for section in sections:
            for i, chunk in enumerate(chunk_text(section["content"])):
                chunks.append({"section": section["section"], "chunk_index": i, "content": chunk})
    else:
        for i, chunk in enumerate(chunk_text(text)):
            chunks.append({"section": None, "chunk_index": i, "content": chunk})
    return {"metadata": extract_metadata(path, text), "sections": sections, "chunks": chunks, "text": text}
